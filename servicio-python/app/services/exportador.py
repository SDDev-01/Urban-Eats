"""Exporta cualquier reporte a CSV, Excel (openpyxl) o PDF (fpdf2)."""
import csv
import io
from datetime import date, datetime

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.excepciones import ReglaDeNegocio

ROJO_MARCA = "A42828"
TIPOS_MIME = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def exportar(titulo: str, columnas: list[tuple[str, str]], filas: list[dict], formato: str) -> tuple[bytes, str, str]:
    """Devuelve (contenido, tipo MIME, nombre de archivo).

    columnas: lista de (clave en la fila, titulo para mostrar).
    """
    nombre_base = titulo.lower().replace(" ", "-")
    marca = datetime.now().strftime("%Y%m%d-%H%M")
    if formato == "csv":
        contenido = _a_csv(columnas, filas)
    elif formato == "xlsx":
        contenido = _a_excel(titulo, columnas, filas)
    elif formato == "pdf":
        contenido = _a_pdf(titulo, columnas, filas)
    else:
        raise ReglaDeNegocio("Formato no soportado. Use json, csv, xlsx o pdf.")
    return contenido, TIPOS_MIME[formato], f"{nombre_base}-{marca}.{formato}"


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float):
        return f"{valor:,.0f}".replace(",", ".") if valor.is_integer() else f"{valor:,.1f}"
    return str(valor)


def _a_csv(columnas, filas) -> bytes:
    salida = io.StringIO()
    # Punto y coma: Excel en espanol lo abre separado en columnas sin configurar nada.
    escritor = csv.writer(salida, delimiter=";")
    escritor.writerow([titulo for _, titulo in columnas])
    for fila in filas:
        escritor.writerow([_texto(fila.get(clave)) for clave, _ in columnas])
    # BOM al inicio para que Excel reconozca las tildes (UTF-8).
    return ("﻿" + salida.getvalue()).encode("utf-8")


def _a_excel(titulo, columnas, filas) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]  # Excel limita el nombre de la hoja a 31 caracteres
    hoja.append([f"Urban Eats – {titulo}"])
    hoja["A1"].font = Font(bold=True, size=14, color=ROJO_MARCA)
    hoja.append([f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    hoja.append([])
    hoja.append([titulo_columna for _, titulo_columna in columnas])
    for celda in hoja[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=ROJO_MARCA)
        celda.alignment = Alignment(horizontal="center")
    for fila in filas:
        hoja.append([fila.get(clave) for clave, _ in columnas])
    for indice, (clave, titulo_columna) in enumerate(columnas, start=1):
        ancho = max([len(titulo_columna)] + [len(_texto(fila.get(clave))) for fila in filas[:200]])
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(ancho + 2, 10), 50)
    hoja.freeze_panes = "A5"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _latin1(texto: str) -> str:
    """Las fuentes basicas del PDF solo tienen caracteres latin-1: se cambian los demas."""
    return texto.replace("–", "-").replace("×", "x").encode("latin-1", "replace").decode("latin-1")


def _a_pdf(titulo, columnas, filas) -> bytes:
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 15)
    pdf.set_text_color(164, 40, 40)
    pdf.cell(0, 9, _latin1(f"Urban Eats - {titulo}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(80, 80, 80)
    pdf.cell(0, 6, _latin1(f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} - {len(filas)} registro(s)"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    ancho_util = pdf.w - pdf.l_margin - pdf.r_margin
    ancho_columna = ancho_util / max(len(columnas), 1)
    pdf.set_fill_color(164, 40, 40)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 8)
    for _, titulo_columna in columnas:
        pdf.cell(ancho_columna, 7, _latin1(titulo_columna)[:30], border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_text_color(17, 20, 23)
    pdf.set_font("Helvetica", "", 7.5)
    for numero, fila in enumerate(filas):
        pdf.set_fill_color(243, 241, 226) if numero % 2 else pdf.set_fill_color(255, 255, 255)
        for clave, _ in columnas:
            valor = _latin1(_texto(fila.get(clave)))
            # Se recorta para que quepa en la celda (el detalle completo esta en Excel).
            while pdf.get_string_width(valor) > ancho_columna - 2 and len(valor) > 3:
                valor = valor[:-4] + "..."
            pdf.cell(ancho_columna, 6, valor, border=1, fill=True)
        pdf.ln()
    return bytes(pdf.output())
